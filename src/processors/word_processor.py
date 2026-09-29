from pathlib import Path

from docx import Document

from src.utils import process_text_units


def extract_word_content(file_path):
    file_path = Path(file_path)

    if file_path.suffix.lower() != ".docx":
        raise ValueError(
            "The provided file is not a DOCX file."
        )

    document = Document(file_path)

    content = []

    for paragraph_number, paragraph in enumerate(
        document.paragraphs,
        start=1
    ):
        text = paragraph.text.strip()

        if text:
            content.append({
                "source_type": "paragraph",
                "source_number": paragraph_number,
                "text": text,
                "text_length": len(text)
            })

    for table_number, table in enumerate(
        document.tables,
        start=1
    ):
        rows = []

        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                rows.append(cells)

        if rows:
            table_text = "\n".join(
                " | ".join(row)
                for row in rows
            )

            content.append({
                "source_type": "table",
                "source_number": table_number,
                "text": table_text,
                "text_length": len(table_text)
            })

    return content


def process_word(file_path):
    file_path = Path(file_path)

    text_units = extract_word_content(file_path)

    yield from process_text_units(
        text_units,
        file_path.name
    )